#!/usr/bin/env python3
"""
Pembuat GOLDEN FILE Keselamatan Listrik (Tahap 6): "sama dengan Excel".

Cara kerja sama dengan buat_golden_pompa.py: tulis isian tiap kasus ke Input 2 (bacaan ESA + klasifikasi + guard band),
LibreOffice menghitung ulang SEMUA rumus, lalu hasil View Report (terukur, U95, batas, Laik/Tidak Laik, kesimpulan)
disimpan sebagai "harapan" di JSON.

Pakai:
  python3 tools/tera/buat_golden_listrik.py <Master_Keselamatan_Listrik_terisi.xlsx> <folder_kerja> [keluaran.json]
Hanya data uji. Berkas Excel TIDAK ikut commit.
"""
import json, os, shutil, subprocess, sys, datetime
import openpyxl


def siapkan_profil(profil):
    """Profil LibreOffice yang MEMAKSA hitung ulang semua rumus (sama dengan buat_golden_pompa.py)."""
    os.makedirs(os.path.join(profil, 'user'), exist_ok=True)
    with open(os.path.join(profil, 'user', 'registrymodifications.xcu'), 'w', encoding='utf-8') as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
                'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
                '<value>0</value></prop></item>\n</oor:items>\n')


def hitung_ulang(masuk_dir, keluar_dir, profil):
    berkas = [os.path.join(masuk_dir, f) for f in sorted(os.listdir(masuk_dir)) if f.endswith('.xlsx')]
    subprocess.run(['soffice', f'-env:UserInstallation=file://{profil}', '--headless', '--convert-to', 'xlsx',
                    '--outdir', keluar_dir, *berkas], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def tabel(ws, kol_titik, kol_koreksi, kol_u, baris):
    out = []
    for r in baris:
        t, c, u = ws[f'{kol_titik}{r}'].value, ws[f'{kol_koreksi}{r}'].value, ws[f'{kol_u}{r}'].value
        if all(isinstance(x, (int, float)) for x in (t, c, u)):
            out.append(dict(titik=t, koreksi=c, u95=u))
    return out

SUMBER, KERJA = sys.argv[1], sys.argv[2]
KELUAR = sys.argv[3] if len(sys.argv) > 3 else 'tests/tera/fixtures/listrik-golden-v1.json'

ID = ['tegangan', 'resistansi_pembumian', 'resistansi_isolasi', 'arus_bocor_peralatan', 'arus_bocor_ap']
K = lambda kelas, ap, pe, metode: dict(kelas=kelas, tipeAP=ap, kabelPE=pe, metodeBocor=metode)

# (nama, bacaan ESA [tegangan, PE, isolasi, bocor peralatan, bocor AP] (None = tidak diisi), klasifikasi, guard band w)
KASUS = [
    ('laik-kelas1-bf-langsung', [221.5, 0.12, 250.0, 45.0, 20.0], K('Kelas I', 'Tipe BF', 'Dapat dilepas - bersama kabel', 'Langsung'), 1),
    ('kelas2-tipe-b-alternatif', [218.0, None, 120.0, 300.0, None], K('Kelas II', 'Tipe B', '', 'Alternatif'), 1),
    ('cf-ap-gagal-diferensial', [230.0, 0.095, 50.0, 480.0, 49.8], K('Kelas I', 'Tipe CF', 'Kabel catu daya saja', 'Diferensial'), 1),
    ('pe-dekat-batas-laik', [220.0, 0.305, 80.0, 60.0, 10.0], K('Kelas I', 'Tipe BF', 'Tidak dapat dilepas', 'Langsung'), 1),
    ('pe-guard-band-gagal', [220.0, 0.308, 80.0, 60.0, 10.0], K('Kelas I', 'Tipe BF', 'Tidak dapat dilepas', 'Langsung'), 1),
    ('pe-sederhana-w0', [220.0, 0.308, 80.0, 60.0, 10.0], K('Kelas I', 'Tipe BF', 'Tidak dapat dilepas', 'Langsung'), 0),
    ('isolasi-gagal', [221.0, 0.1, 2.03, 40.0, None], K('Kelas I', 'Tidak ada', 'Instalasi permanen', 'Langsung'), 1),
    ('bocor-gagal-kelas2-langsung', [221.0, None, 300.0, 150.0, 30.0], K('Kelas II', 'Tipe BF', '', 'Langsung'), 1),
    ('tegangan-luar-rentang', [245.0, 0.2, 500.0, 100.0, 4000.0], K('Kelas I', 'Tipe BF', 'Dapat dilepas - tanpa kabel', 'Alternatif'), 1),
    ('kotak-multipel-alternatif', [215.0, 0.45, 15.0, 900.0, None], K('Kelas I', 'Tipe B', 'Kotak kontak multipel', 'Alternatif'), 1),
    ('bocor-kecil-proporsional', [205.0, 0.05, 9.0, 5.0, 3.0], K('Kelas I', 'Tipe CF', 'Dapat dilepas - bersama kabel', 'Langsung'), 1),
    ('pe-dapat-dilepas-tanpa-kabel-gagal', [220.0, 0.22, 100.0, 50.0, 10.0], K('Kelas I', 'Tipe BF', 'Dapat dilepas - tanpa kabel', 'Langsung'), 1),
]


def tulis_kasus(k, jalur):
    nama, bacaan, kl, w = k
    wb = openpyxl.load_workbook(SUMBER)
    ws = wb['Input 2']
    for i, v in enumerate(bacaan):
        ws[f'H{14 + i}'] = v
    ws['M13'], ws['M14'], ws['M15'], ws['M16'], ws['M17'] = kl['kelas'], kl['tipeAP'], kl['kabelPE'] or None, kl['metodeBocor'], w
    wb.save(jalur)


def angka(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def main():
    masuk, keluar, profil = [os.path.join(KERJA, d) for d in ('masuk', 'keluar', 'profil')]
    for d in (masuk, keluar):
        shutil.rmtree(d, ignore_errors=True)
        os.makedirs(d)
    siapkan_profil(profil)
    for n, k in enumerate(KASUS):
        tulis_kasus(k, os.path.join(masuk, f'kasus{n:02d}.xlsx'))
    hitung_ulang(masuk, keluar, profil)
    hasil = []
    for n, (nama, bacaan, kl, w) in enumerate(KASUS):
        wb = openpyxl.load_workbook(os.path.join(keluar, f'kasus{n:02d}.xlsx'), data_only=True)
        vr, cal = wb['View Report'], wb['Calibrator 1']
        harapan = {}
        for i, id_ in enumerate(ID):
            r = 11 + i
            harapan[id_] = dict(terukur=angka(vr[f'H{r}'].value), u95=angka(vr[f'X{r}'].value), hasil=vr[f'L{r}'].value,
                                batas=angka(vr[f'P{r}'].value) if r > 11 else 22, toleransi=vr[f'I{r}'].value)
        hasil.append(dict(
            nama=nama, guardBand=w,
            data=dict(klasifikasi={k: v for k, v in kl.items() if v}, listrik={id_: v for id_, v in zip(ID, bacaan) if v is not None}),
            kalibrator=dict(listrik=dict(
                tegangan=tabel(cal, 'R', 'S', 'T', range(28, 34)), resistansi=tabel(cal, 'R', 'S', 'T', range(37, 41)),
                isolasi=tabel(cal, 'R', 'S', 'T', range(44, 46)), arus_bocor=tabel(cal, 'R', 'S', 'T', range(49, 53)))),
            harapan=dict(listrik=harapan, kesimpulan=vr['C18'].value),
        ))
    os.makedirs(os.path.dirname(KELUAR) or '.', exist_ok=True)
    with open(KELUAR, 'w', encoding='utf-8') as f:
        json.dump(dict(sumber='Master_Keselamatan_Listrik_terisi.xlsx (26 Okt 2026), dihitung ulang penuh dengan LibreOffice',
                       dibuat=datetime.date.today().isoformat(), kasus=hasil), f, ensure_ascii=False, indent=1)
    print('selesai:', KELUAR, len(hasil), 'kasus')


main()
