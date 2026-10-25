#!/usr/bin/env python3
"""
Pembuat GOLDEN FILE Infusion Pump (Tahap 6): "sama dengan Excel".

Cara kerja (perumpamaan: menimbang beban yang sama di dua timbangan lalu membandingkan):
  1. Ambil Excel Infusion Pump yang SUDAH TERISI (berisi tabel kalibrator IDA + ESA).
  2. Untuk tiap kasus uji, tulis angka isian ke Input 1 / Input 2, simpan sebagai berkas baru.
  3. LibreOffice menghitung ulang SEMUA rumus (profil khusus OOXMLRecalcMode=0 supaya tidak memakai nilai lama).
  4. Baca hasil Excel (Ketidakpastian + View Report) -> disimpan sebagai "harapan" di JSON.
Mesin hitung di aplikasi (src/tera/shared/hitungPompa.ts) lalu diuji agar menghasilkan angka yang sama.

Pakai:
  python3 tools/tera/buat_golden_pompa.py <Master_Infusion_Pump_terisi.xlsx> <folder_kerja> [keluaran.json]
Hanya data uji (BO 999 / DATA UJI COBA). Berkas Excel TIDAK ikut commit (_referensi/ & folder kerja).
"""
import json, os, shutil, subprocess, sys, datetime
import openpyxl

SUMBER = sys.argv[1]
KERJA = sys.argv[2]
KELUAR = sys.argv[3] if len(sys.argv) > 3 else 'tests/tera/fixtures/infusion-golden-v1.json'

# ───────── kasus uji ─────────
# aliran: setting -> 6 bacaan standar IDA (mentah, sebelum koreksi)
KASUS = [
    dict(nama='laik-lengkap', skala='Digital', resolusi=0.1,
         listrik=dict(tegangan=221.5, resistansi_pembumian=0.12, resistansi_isolasi=250.0, arus_bocor=45.0),
         aliran={25: [25.12, 24.95, 25.2, 25.05, 24.98, 25.1], 50: [50.3, 49.85, 50.15, 50.4, 49.9, 50.2],
                 100: [100.4, 99.6, 100.8, 100.2, 99.9, 100.5], 150: [150.5, 149.7, 150.9, 150.2, 149.8, 150.4],
                 200: [201.0, 199.5, 200.8, 200.3, 199.8, 200.6]},
         sumbat=(100, [9.8, 10.1, 9.9, 10.3, 10.0, 9.7])),
    dict(nama='satu-titik-aliran-gagal-masih-70persen', skala='Digital', resolusi=0.1,
         listrik=dict(tegangan=219.0, resistansi_pembumian=0.2, resistansi_isolasi=100.0, arus_bocor=120.0),
         aliran={25: [28.4, 28.1, 28.6, 28.3, 28.2, 28.5], 50: [50.3, 49.85, 50.15, 50.4, 49.9, 50.2],
                 100: [100.4, 99.6, 100.8, 100.2, 99.9, 100.5], 150: [150.5, 149.7, 150.9, 150.2, 149.8, 150.4],
                 200: [201.0, 199.5, 200.8, 200.3, 199.8, 200.6]},
         sumbat=(100, [12.0, 12.4, 11.8, 12.2, 12.1, 12.3])),
    dict(nama='tiga-titik-aliran-gagal', skala='Digital', resolusi=0.1,
         listrik=dict(tegangan=221.5, resistansi_pembumian=0.12, resistansi_isolasi=250.0, arus_bocor=45.0),
         aliran={25: [28.4, 28.1, 28.6, 28.3, 28.2, 28.5], 50: [56.0, 55.6, 56.3, 55.9, 56.1, 55.8],
                 100: [100.4, 99.6, 100.8, 100.2, 99.9, 100.5], 150: [166.0, 165.5, 166.4, 165.8, 166.2, 165.9],
                 200: [201.0, 199.5, 200.8, 200.3, 199.8, 200.6]},
         sumbat=(100, [9.8, 10.1, 9.9, 10.3, 10.0, 9.7])),
    dict(nama='listrik-gagal', skala='Digital', resolusi=0.1,
         listrik=dict(tegangan=221.5, resistansi_pembumian=0.35, resistansi_isolasi=1.5, arus_bocor=620.0),
         aliran={25: [25.12, 24.95, 25.2, 25.05, 24.98, 25.1], 50: [50.3, 49.85, 50.15, 50.4, 49.9, 50.2],
                 100: [100.4, 99.6, 100.8, 100.2, 99.9, 100.5], 150: [150.5, 149.7, 150.9, 150.2, 149.8, 150.4],
                 200: [201.0, 199.5, 200.8, 200.3, 199.8, 200.6]},
         sumbat=(100, [9.8, 10.1, 9.9, 10.3, 10.0, 9.7])),
    dict(nama='penyumbatan-gagal', skala='Digital', resolusi=0.1,
         listrik=dict(tegangan=221.5, resistansi_pembumian=0.12, resistansi_isolasi=250.0, arus_bocor=45.0),
         aliran={25: [25.12, 24.95, 25.2, 25.05, 24.98, 25.1], 50: [50.3, 49.85, 50.15, 50.4, 49.9, 50.2],
                 100: [100.4, 99.6, 100.8, 100.2, 99.9, 100.5], 150: [150.5, 149.7, 150.9, 150.2, 149.8, 150.4],
                 200: [201.0, 199.5, 200.8, 200.3, 199.8, 200.6]},
         sumbat=(100, [24.5, 24.9, 24.2, 25.1, 24.7, 24.4])),
    dict(nama='analog-resolusi-1', skala='Analog', resolusi=1,
         listrik=dict(tegangan=225.0, resistansi_pembumian=0.1, resistansi_isolasi=500.0, arus_bocor=80.0),
         aliran={25: [25.5, 24.6, 25.8, 25.2, 24.9, 25.4], 50: [51.0, 49.2, 50.8, 50.5, 49.6, 50.9],
                 100: [101.0, 99.0, 100.5, 100.0, 99.5, 101.5], 150: [151.0, 149.0, 150.5, 150.0, 149.5, 151.5],
                 200: [202.0, 199.0, 201.0, 200.0, 199.5, 202.5]},
         sumbat=(100, [10.0, 10.5, 9.5, 10.0, 10.5, 9.5])),
    # tepi tabel kalibrator: di bawah titik pertama, tepat di titik, di atas titik terakhir
    dict(nama='tepi-tabel-kalibrator', skala='Digital', resolusi=0.1,
         listrik=dict(tegangan=99.0, resistansi_pembumian=0.05, resistansi_isolasi=0.5, arus_bocor=5.0),
         aliran={5: [5.0, 4.9, 5.1, 5.0, 4.95, 5.05], 10.31: [10.31, 10.3, 10.32, 10.31, 10.29, 10.33],
                 100: [100.4, 99.6, 100.8, 100.2, 99.9, 100.5], 294.83: [294.83, 294.8, 294.9, 294.85, 294.8, 294.86],
                 300: [300.5, 299.5, 300.8, 300.2, 299.9, 300.4]},
         sumbat=(100, [0.0, 0.1, 0.0, 0.1, 0.0, 0.1])),
    dict(nama='listrik-atas-tabel', skala='Digital', resolusi=0.1,
         listrik=dict(tegangan=250.0, resistansi_pembumian=1.5, resistansi_isolasi=15.0, arus_bocor=520.0),
         aliran={25: [25.12, 24.95, 25.2, 25.05, 24.98, 25.1], 50: [50.3, 49.85, 50.15, 50.4, 49.9, 50.2],
                 100: [100.4, 99.6, 100.8, 100.2, 99.9, 100.5], 150: [150.5, 149.7, 150.9, 150.2, 149.8, 150.4],
                 200: [201.0, 199.5, 200.8, 200.3, 199.8, 200.6]},
         sumbat=(100, [30.0, 31.0, 29.0, 30.5, 29.5, 30.0])),
]

FISIK_ID = ['badan', 'kontrol', 'kabel']
# 6 kotak: (fisik, fungsi) per butir
FISIK_KASUS = {
    'laik-lengkap': [('Baik', 'Berfungsi')] * 3,
    'satu-titik-aliran-gagal-masih-70persen': [('Baik', 'Berfungsi'), ('Baik', 'Berfungsi'), ('Baik', 'Tidak Berfungsi')],
    'tiga-titik-aliran-gagal': [('Baik', 'Berfungsi')] * 3,
    'listrik-gagal': [('Baik', 'Berfungsi')] * 3,
    'penyumbatan-gagal': [('Baik', 'Berfungsi')] * 3,
    'analog-resolusi-1': [('Baik', 'Berfungsi'), ('Tidak Baik', 'Tidak Berfungsi'), ('Baik', 'Berfungsi')],
    'tepi-tabel-kalibrator': [('Baik', 'Berfungsi'), ('Baik', 'Tidak Berfungsi'), ('Tidak Baik', 'Tidak Berfungsi')],
    'listrik-atas-tabel': [('Baik', 'Berfungsi')] * 3,
}

LISTRIK_BARIS = dict(tegangan=(13, 10), resistansi_pembumian=(14, 11), resistansi_isolasi=(15, 12), arus_bocor=(16, 13))
KOLOM_RUN = 'EFGHIJ'


def tulis_kasus(k, jalur):
    wb = openpyxl.load_workbook(SUMBER)
    i1, i2 = wb['Input 1'], wb['Input 2']
    i1['E19'] = k['resolusi']
    i1['E20'] = k['skala']
    for id_, (r_in, _) in LISTRIK_BARIS.items():
        i2[f'H{r_in}'] = k['listrik'][id_]
    for baris, (setting, run) in zip(range(22, 27), k['aliran'].items()):
        i2[f'C{baris}'] = setting
        for kol, v in zip(KOLOM_RUN, run):
            i2[f'{kol}{baris}'] = v
    i2['C32'] = k['sumbat'][0]
    for kol, v in zip(KOLOM_RUN, k['sumbat'][1]):
        i2[f'{kol}32'] = v
    wb.save(jalur)


def hitung_ulang(masuk_dir, keluar_dir, profil):
    os.makedirs(keluar_dir, exist_ok=True)
    berkas = [os.path.join(masuk_dir, f) for f in sorted(os.listdir(masuk_dir)) if f.endswith('.xlsx')]
    subprocess.run(['soffice', f'-env:UserInstallation=file://{profil}', '--headless', '--convert-to', 'xlsx',
                    '--outdir', keluar_dir, *berkas], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def siapkan_profil(profil):
    """Profil LibreOffice yang MEMAKSA hitung ulang semua rumus (bukan memakai nilai tersimpan)."""
    os.makedirs(os.path.join(profil, 'user'), exist_ok=True)
    xcu = os.path.join(profil, 'user', 'registrymodifications.xcu')
    with open(xcu, 'w', encoding='utf-8') as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
                'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
                '<value>0</value></prop></item>\n</oor:items>\n')


def tabel(ws, kol_titik, kol_koreksi, kol_u, baris):
    out = []
    for r in baris:
        t, c, u = ws[f'{kol_titik}{r}'].value, ws[f'{kol_koreksi}{r}'].value, ws[f'{kol_u}{r}'].value
        if isinstance(t, (int, float)) and isinstance(c, (int, float)) and isinstance(u, (int, float)):
            out.append(dict(titik=t, koreksi=c, u95=u))
    return out


def baca_kalibrator(wb):
    k = wb['Calibrator 1']
    return dict(
        aliran=tabel(k, 'A', 'C', 'D', range(21, 26)),
        tekanan=tabel(k, 'G', 'I', 'J', range(21, 27)),
        listrik=dict(
            tegangan=tabel(k, 'R', 'S', 'T', range(61, 67)),
            resistansi=tabel(k, 'R', 'S', 'T', range(70, 74)),
            isolasi=tabel(k, 'R', 'S', 'T', range(77, 79)),
            arus_bocor=tabel(k, 'R', 'S', 'T', range(82, 86)),
        ),
    )


def vonis(x):
    return x if x in ('Laik', 'Tidak Laik') else '-'


def baca_harapan(wb):
    kt, vr = wb['Ketidakpastian'], wb['View Report']
    listrik = {}
    for id_, (_, r) in LISTRIK_BARIS.items():
        listrik[id_] = dict(terukur=vr[f'H{r}'].value, u95=vr[f'X{r}'].value, hasil=vonis(vr[f'L{r}'].value))
    aliran = []
    for r in range(24, 29):
        aliran.append(dict(setting=kt[f'D{r}'].value, rataRata=kt[f'K{r}'].value, koreksi=kt[f'L{r}'].value,
                           uTipeA=kt[f'N{r}'].value, uSertifikat=kt[f'Q{r}'].value, uResolusi=kt['T24'].value,
                           uDrift=kt[f'W{r}'].value, uc=kt[f'Y{r}'].value, vEff=kt[f'Z{r}'].value, k=kt[f'AA{r}'].value,
                           u95=kt[f'AB{r}'].value, hasil=vonis(vr[f'L{r - 6}'].value)))
    sumbat = dict(rataRata=kt['K37'].value, uTipeA=kt['N37'].value, uSertifikat=kt['Q37'].value, uc=kt['Y37'].value,
                  vEff=kt['Z37'].value, k=kt['AA37'].value, u95=kt['AB37'].value, hasil=vonis(vr['L32'].value))
    return dict(listrik=listrik, aliran=aliran, sumbat=sumbat, nilaiListrik=vr['AE3'].value, nilaiKinerja=vr['AE4'].value)


def main():
    masuk, keluar, profil = [os.path.join(KERJA, d) for d in ('masuk', 'keluar', 'profil')]
    for d in (masuk, keluar):
        shutil.rmtree(d, ignore_errors=True)
        os.makedirs(d)
    siapkan_profil(profil)
    for n, k in enumerate(KASUS):
        tulis_kasus(k, os.path.join(masuk, f'kasus{n:02d}.xlsx'))
    hitung_ulang(masuk, keluar, profil)
    hasil = []
    for n, k in enumerate(KASUS):
        wb = openpyxl.load_workbook(os.path.join(keluar, f'kasus{n:02d}.xlsx'), data_only=True)
        fisik = {i: dict(fisik=a, fungsi=b) for i, (a, b) in zip(FISIK_ID, FISIK_KASUS[k['nama']])}
        hasil.append(dict(
            nama=k['nama'],
            titikAliran=list(k['aliran'].keys()),
            data=dict(identitas=dict(resolusi=k['resolusi'], jenisSkala=k['skala']), fisik=fisik, listrik=k['listrik'],
                      aliran={str(s): v for s, v in k['aliran'].items()},
                      sumbat=dict(setting=k['sumbat'][0], bacaan=k['sumbat'][1])),
            kalibrator=baca_kalibrator(wb),
            harapan=baca_harapan(wb),
        ))
    os.makedirs(os.path.dirname(KELUAR) or '.', exist_ok=True)
    with open(KELUAR, 'w', encoding='utf-8') as f:
        json.dump(dict(sumber='Master_Infusion_Pump_terisi.xlsx (26 Okt 2026), dihitung ulang penuh dengan LibreOffice',
                       dibuat=datetime.date.today().isoformat(),
                       catatan='Fisik & fungsi memakai aturan baru (semua kotak); harapan Excel hanya untuk listrik, aliran, penyumbatan.',
                       kasus=hasil), f, ensure_ascii=False, indent=1)
    print('selesai:', KELUAR, len(hasil), 'kasus')


main()
