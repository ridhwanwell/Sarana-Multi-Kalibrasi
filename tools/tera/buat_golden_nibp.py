#!/usr/bin/env python3
"""
Pembuat GOLDEN FILE NIBP (Tahap 6): "sama dengan Excel".
Tulis bacaan alat tiap kasus ke Input 2!D24:I35 (+ resolusi Input 1!E19), kosongkan uji listrik (alat standar hanya
NIBP Simulator + termohigrometer), LibreOffice menghitung ulang SEMUA rumus, lalu hasil Ketidakpastian!J31:AE42 +
View Report!L19:L30, AE2:AE5, C32 disimpan sebagai "harapan" di JSON.
Pakai: python3 tools/tera/buat_golden_nibp.py <Master_NIBP_terisi.xlsx> <folder_kerja> [keluaran.json]
"""
import json, os, shutil, subprocess, sys, datetime
import openpyxl

SUMBER, KERJA = sys.argv[1], sys.argv[2]
KELUAR = sys.argv[3] if len(sys.argv) > 3 else 'tests/tera/fixtures/nibp-golden-v1.json'
SETTING = [60, 40, 30, 120, 93, 80, 150, 117, 100, 200, 167, 150]
ID = [f's{i // 3 + 1}_{["systole", "map", "diastole"][i % 3]}' for i in range(12)]


def r6(base, d):  # 6 bacaan: setting + base + pola
    return [base + x for x in d]


POLA = [0, 1, -1, 1, 0, -1]
KASUS = [
    # nama, resolusi, bacaan per titik (12 × list), 
    ('laik-contoh-excel', 1, [[60, 61, 59, 61, 60, 59], [41, 40, 41, 39, 40, 41], [30, 29, 31, 30, 31, 29], [120, 121, 119, 121, 120, 119], [94, 93, 94, 92, 93, 94], [80, 79, 81, 80, 81, 79],
                              [150, 151, 149, 151, 150, 149], [118, 117, 118, 116, 117, 118], [100, 99, 101, 100, 101, 99], [200, 201, 199, 201, 200, 199], [168, 167, 168, 166, 167, 168], [150, 149, 151, 150, 151, 149]]),
    ('semua-tepat', 1, [[s] * 6 for s in SETTING]),
    ('bias-plus-5', 1, [[s + 5 + x for x in POLA] for s in SETTING]),
    ('bias-plus-7-gagal', 1, [[s + 7 + x for x in POLA] for s in SETTING]),
    ('empat-titik-gagal-kurang-70persen', 1, [[s + (10 if i in (0, 3, 6, 9) else 0) + x for x in POLA] for i, s in enumerate(SETTING)]),
    ('tiga-titik-gagal-masih-75persen', 1, [[s + (10 if i in (0, 3, 6) else 0) + x for x in POLA] for i, s in enumerate(SETTING)]),
    ('empat-bacaan-saja', 1, [[s + x for x in POLA[:4]] for s in SETTING]),
    ('lima-bacaan-resolusi-01', 0.1, [[s + x * 0.4 for x in [0.5, -0.3, 0.2, 0.1, -0.4]] for s in SETTING]),
    ('sebar-besar', 1, [[s + x for x in [4, -4, 3, -3, 2, -2]] for s in SETTING]),
]


def siapkan_profil(profil):
    os.makedirs(os.path.join(profil, 'user'), exist_ok=True)
    with open(os.path.join(profil, 'user', 'registrymodifications.xcu'), 'w', encoding='utf-8') as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
                'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
                '<value>0</value></prop></item>\n</oor:items>\n')


def hitung_ulang(masuk_dir, keluar_dir, profil):
    berkas = [os.path.join(masuk_dir, f) for f in sorted(os.listdir(masuk_dir)) if f.endswith('.xlsx')]
    subprocess.run(['soffice', f'-env:UserInstallation=file://{profil}', '--headless', '--convert-to', 'xlsx',
                    '--outdir', keluar_dir, *berkas], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def tulis(k, jalur):
    nama, res, bacaan = k
    wb = openpyxl.load_workbook(SUMBER)
    wb['Input 1']['E19'] = res
    ws = wb['Input 2']
    for r in range(13, 18):
        ws[f'H{r}'] = None  # tanpa uji listrik
    for i, run in enumerate(bacaan):
        for j, col in enumerate('DEFGHI'):
            ws[f'{col}{24 + i}'] = run[j] if j < len(run) else None
    wb.save(jalur)


def angka(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def main():
    masuk, keluar, profil = [os.path.join(KERJA, d) for d in ('masuk', 'keluar', 'profil')]
    for d in (masuk, keluar):
        shutil.rmtree(d, ignore_errors=True)
        os.makedirs(d)
    siapkan_profil(profil)
    for n, k in enumerate(KASUS):
        tulis(k, os.path.join(masuk, f'kasus{n:02d}.xlsx'))
    hitung_ulang(masuk, keluar, profil)
    hasil = []
    for n, (nama, res, bacaan) in enumerate(KASUS):
        wb = openpyxl.load_workbook(os.path.join(keluar, f'kasus{n:02d}.xlsx'), data_only=True)
        kt, vr, cal = wb['Ketidakpastian'], wb['View Report'], wb['Calibrator 1']
        titik = []
        for i in range(12):
            r = 31 + i
            titik.append(dict(id=ID[i], setting=kt[f'C{r}'].value, rataRata=angka(kt[f'J{r}'].value), koreksi=angka(kt[f'K{r}'].value),
                              uTipeA=angka(kt[f'M{r}'].value), uSertifikat=angka(kt[f'P{r}'].value), uResolusi=angka(kt['S31'].value),
                              uDrift=angka(kt[f'V{r}'].value), uc=angka(kt[f'X{r}'].value), vEff=angka(kt[f'Y{r}'].value), k=angka(kt[f'Z{r}'].value),
                              u95=angka(kt[f'AA{r}'].value), koreksiStd=angka(kt[f'AD{r}'].value), uStd=angka(kt[f'AE{r}'].value),
                              hasil=vr[f'L{19 + i}'].value, cetakAlat=angka(vr[f'E{19 + i}'].value)))
        tabel = []
        for r in range(21, 33):
            a, c, d = cal[f'A{r}'].value, cal[f'C{r}'].value, cal[f'D{r}'].value
            if all(isinstance(x, (int, float)) for x in (a, c, d)):
                tabel.append(dict(titik=a, koreksi=c, u95=d))
        hasil.append(dict(nama=nama, data=dict(identitas=dict(resolusi=res), nibp={ID[i]: run for i, run in enumerate(bacaan)}),
                          kalibrator=dict(nibp=tabel),
                          harapan=dict(titik=titik, skor=dict(fisik=vr['AE2'].value, listrik=vr['AE3'].value, kinerja=vr['AE4'].value, total=vr['AE5'].value),
                                       kesimpulan=vr['C32'].value)))
    with open(KELUAR, 'w', encoding='utf-8') as f:
        json.dump(dict(sumber='Master_NIBP_terisi.xlsx (26 Okt 2026), dihitung ulang penuh dengan LibreOffice; uji listrik dikosongkan',
                       dibuat=datetime.date.today().isoformat(), kasus=hasil), f, ensure_ascii=False, indent=1)
    print('selesai:', KELUAR, len(hasil), 'kasus')


main()
